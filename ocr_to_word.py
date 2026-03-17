from __future__ import annotations

import argparse
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse
from urllib.error import URLError
from urllib.request import urlopen

import pytesseract
from docx import Document
from docx.shared import Pt, RGBColor
from PIL import Image, ImageStat
from pytesseract import Output


@dataclass
class WordBox:
    text: str
    x: int
    y: int
    w: int
    h: int
    r: int
    g: int
    b: int

    @property
    def right(self) -> int:
        return self.x + self.w

    @property
    def cy(self) -> float:
        return self.y + (self.h / 2)


def _is_url(path_or_url: str) -> bool:
    parsed = urlparse(path_or_url)
    return parsed.scheme in {"http", "https"} and bool(parsed.netloc)


def _open_image(source: str) -> Image.Image:
    if _is_url(source):
        with urlopen(source) as response:
            data = response.read()
        return Image.open(BytesIO(data)).convert("RGB")
    return Image.open(source).convert("RGB")


def _sample_color(img: Image.Image, x: int, y: int, w: int, h: int) -> tuple[int, int, int]:
    left = max(0, x)
    top = max(0, y)
    right = min(img.width, x + w)
    bottom = min(img.height, y + h)
    if left >= right or top >= bottom:
        return (0, 0, 0)
    stat = ImageStat.Stat(img.crop((left, top, right, bottom)))
    median = [int(value) for value in stat.median[:3]]
    return (median[0], median[1], median[2])


def _extract_word_boxes(img: Image.Image) -> list[WordBox]:
    data = pytesseract.image_to_data(img, output_type=Output.DICT, config="--oem 3 --psm 6")
    words: list[WordBox] = []
    total = len(data["text"])
    for i in range(total):
        text = data["text"][i].strip()
        if not text:
            continue
        try:
            confidence = float(data["conf"][i])
        except ValueError:
            confidence = -1
        if confidence <= 0:
            continue
        x = int(data["left"][i])
        y = int(data["top"][i])
        w = int(data["width"][i])
        h = int(data["height"][i])
        r, g, b = _sample_color(img, x, y, w, h)
        words.append(WordBox(text=text, x=x, y=y, w=w, h=h, r=r, g=g, b=b))
    return words


def _group_lines(words: list[WordBox]) -> list[list[WordBox]]:
    if not words:
        return []
    words = sorted(words, key=lambda w: (w.cy, w.x))
    heights = sorted(word.h for word in words)
    median_h = heights[len(heights) // 2]
    y_tolerance = max(10, int(median_h * 0.55))
    lines: list[list[WordBox]] = []
    line_centers: list[float] = []
    for word in words:
        placed = False
        for index, center in enumerate(line_centers):
            if abs(word.cy - center) <= y_tolerance:
                lines[index].append(word)
                line_centers[index] = sum(w.cy for w in lines[index]) / len(lines[index])
                placed = True
                break
        if not placed:
            lines.append([word])
            line_centers.append(word.cy)
    for line in lines:
        line.sort(key=lambda w: w.x)
    lines.sort(key=lambda line: min(word.y for word in line))
    return lines


def convert_image_to_word(source: str, output_docx: str) -> str:
    image = _open_image(source)
    words = _extract_word_boxes(image)
    lines = _group_lines(words)

    document = Document()
    section = document.sections[0]
    usable_width = section.page_width - section.left_margin - section.right_margin
    previous_line_y: float | None = None
    median_line_height = max(14, int(sum(word.h for word in words) / max(1, len(words))))

    for line in lines:
        top = min(word.y for word in line)
        if previous_line_y is not None and top - previous_line_y > int(median_line_height * 1.8):
            document.add_paragraph("")
        previous_line_y = top

        paragraph = document.add_paragraph()
        first_word = line[0]
        paragraph.paragraph_format.left_indent = int((first_word.x / image.width) * usable_width)

        previous_word = None
        for word in line:
            if previous_word is not None:
                gap = max(0, word.x - previous_word.right)
                normalized_gap = max(1, int(round(gap / max(6, previous_word.h * 0.45))))
                paragraph.add_run(" " * normalized_gap)
            run = paragraph.add_run(word.text)
            run.font.color.rgb = RGBColor(word.r, word.g, word.b)
            run.font.size = Pt(max(8, min(22, int(round(word.h * 0.65)))))
            previous_word = word

    output_path = Path(output_docx)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return str(output_path.resolve())


def _default_output_name(source: str) -> str:
    if _is_url(source):
        candidate = Path(urlparse(source).path).name or "output"
    else:
        candidate = Path(source).name
    stem = Path(candidate).stem or "output"
    return f"{stem}.docx"


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Convert an image (local path or URL) into an editable Word file using OCR."
    )
    parser.add_argument("source", help="Image path or URL")
    parser.add_argument(
        "-o",
        "--output",
        default=None,
        help="Output .docx path (defaults to <input-name>.docx in current directory)",
    )
    args = parser.parse_args()

    output_path = args.output or _default_output_name(args.source)
    try:
        saved = convert_image_to_word(args.source, output_path)
    except URLError as exc:
        raise SystemExit(f"Could not download image from URL: {exc}") from exc
    except FileNotFoundError as exc:
        raise SystemExit(f"Image file not found: {exc}") from exc
    except pytesseract.TesseractNotFoundError as exc:
        raise SystemExit(
            "Tesseract OCR is not installed or not in PATH. "
            "Install `tesseract-ocr` and run the command again."
        ) from exc
    print(f"Created editable Word file: {saved}")


if __name__ == "__main__":
    main()
